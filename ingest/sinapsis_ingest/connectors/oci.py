"""Autorizaciones de actividad privada tras el cese de altos cargos.

    ex alto cargo (Person) --Occupancy--> puesto (Position)     hasta el cese
    ex alto cargo (Person) --UnknownLink--> actividad autorizada

Es la fuente que AFIRMA una puerta giratoria (spec §12): la Oficina de
Conflictos de Intereses autoriza —o no— a quien deja un alto cargo a trabajar
en el sector privado en los dos años siguientes (Ley 3/2015, art. 15), y el
Portal de Transparencia publica las autorizaciones. Que «X, ex ministra de
Empleo, fue autorizada a ser consejera de Y» lo dice un documento oficial.

Lo que NO dice, y por eso no se escribe en ninguna parte: que la persona
llegara a ocupar el puesto, ni nada sobre su actuación en el cargo. Una
autorización es una autorización.

## La forma

La verificó el reconocimiento (`docs/fuentes/oci-reconocimiento.md`): una
tabla HTML con estas columnas, en este orden, y fechas `aaaa/mm/dd`:

    Nombre | Alto Cargo | Ministerio | Fecha de Cese |
    Empresa/Actividad autorizada | Fecha de autorización

El nombre viene como «APELLIDOS, NOMBRE» en mayúsculas, y no siempre bien:
«LORA-TAMAYO, D'OCON» no trae nombre de pila. Por eso se publica en el orden
en que viene y no se reordena: darle la vuelta inventaría un nombre.
"""

from __future__ import annotations

import io
import json
import os
import re
import time
import unicodedata
import warnings
import zipfile
from collections.abc import Iterator
from datetime import UTC, date, datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

import httpx
import structlog

from sinapsis_ingest.connectors.base import ParsedRecord, RawDocument
from sinapsis_ingest.connectors.boe import clave
from sinapsis_ingest.normalizado import AristaNormalizada, EntidadNormalizada, Normalizado

log = structlog.get_logger()

BASE = "https://transparencia.gob.es"
#: El buscador del portal con TODAS las autorizaciones —vigentes e
#: históricas: 644 el 25/9/2026—. Enlaza a la exportación en hoja de cálculo.
BUSCADOR = (
    f"{BASE}/servicios-buscador/buscar.htm?categoria=autorizaciones_ind"
    "&lang=es&orderBy=fechaAutorizacion&or=DESC"
)
_EXPORTACION = re.compile(r'href="([^"]*expTab\.htm[^"]*)"')

COLUMNAS = (
    "nombre",
    "alto cargo",
    "ministerio",
    "fecha de cese",
    "empresa/actividad autorizada",
    "fecha de autorizacion",
)


def _plano(texto: str) -> str:
    sin = "".join(
        c for c in unicodedata.normalize("NFKD", texto or "") if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"\s+", " ", sin.lower()).strip()


class _Tablas(HTMLParser):
    """Las tablas de una página, con el enlace de la primera celda."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tablas: list[list[list[str]]] = []
        self.enlaces: list[list[str]] = []
        self._fila: list[str] | None = None
        self._celda: list[str] | None = None
        self._enlace_fila = ""

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "table":
            self.tablas.append([])
            self.enlaces.append([])
        elif tag == "tr" and self.tablas:
            self._fila, self._enlace_fila = [], ""
        elif tag in ("td", "th") and self._fila is not None:
            self._celda = []
        elif tag == "a" and self._celda is not None and not self._enlace_fila:
            self._enlace_fila = dict(attrs).get("href") or ""

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._celda is not None and self._fila is not None:
            self._fila.append(" ".join("".join(self._celda).split()))
            self._celda = None
        elif tag == "tr" and self._fila is not None and self.tablas:
            if self._fila:
                self.tablas[-1].append(self._fila)
                self.enlaces[-1].append(self._enlace_fila)
            self._fila = None

    def handle_data(self, data: str) -> None:
        if self._celda is not None:
            self._celda.append(data)


def _fecha(texto: str) -> date | None:
    """Las dos formas que usa la fuente, sin confundirlas.

    Las páginas de seguimiento escriben «2018/06/01» (año delante) y el
    buscador «24/09/2025» (día delante). Con el año de cuatro cifras en un
    extremo no hay ambigüedad; cualquier otra cosa no se interpreta.
    """
    t = (texto or "").strip()
    m = re.fullmatch(r"(\d{4})/(\d{1,2})/(\d{1,2})", t)
    if m:
        anio, mes, dia = m.groups()
    else:
        m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", t)
        if not m:
            return None
        dia, mes, anio = m.groups()
    try:
        return date(int(anio), int(mes), int(dia))
    except ValueError:
        return None


_PARTICULAS = {"de", "del", "la", "las", "los", "y", "i"}
_TRAS_SEPARADOR = re.compile(r"(^|[-'\u2019\u00b4])(\w)")


def en_minusculas(texto: str) -> str:
    """«ROSA CORDON, RUFINO DE LA» → «Rosa Cordon, Rufino de la».

    Sólo la caja: ni se reordena ni se añade una tilde que la fuente no puso.
    """
    palabras = []
    for p in texto.split(" "):
        bajo = p.lower()
        if bajo in _PARTICULAS:
            palabras.append(bajo)
            continue
        # Mayúscula tras guion y apóstrofo —recto, curvo o el acento agudo
        # suelto que usa la fuente—: «Lora-Tamayo», «D'Ocon».
        palabras.append(_TRAS_SEPARADOR.sub(lambda m: m.group(1) + m.group(2).upper(), bajo))
    return " ".join(palabras)


def nombre_para_cruzar(nombre_fuente: str) -> str:
    """«BAÑEZ GARCIA, FATIMA» → «fatima banez garcia», normalizado.

    SÓLO para buscar coincidencias con el BOE, nunca para enseñarlo: si la
    fuente trae mal el nombre, el reordenado también está mal, y lo único que
    pasa es que no coincide con nada.
    """
    if "," not in nombre_fuente:
        return _plano(nombre_fuente)
    apellidos, nombre = (x.strip() for x in nombre_fuente.split(",", 1))
    # «ROSA CORDON, RUFINO DE LA»: las partículas del final son del apellido.
    partes = nombre.split()
    cola = []
    while partes and partes[-1].lower() in _PARTICULAS:
        cola.insert(0, partes.pop())
    return _plano(" ".join([*partes, *cola, apellidos]))


def _cabecera(celda: str) -> str:
    """Una cabecera, como para compararla.

    El buscador del portal escribe «Empresa / Atividad autorizada» —espacios
    alrededor de la barra, y la errata tal cual—, y las páginas de
    seguimiento «Empresa/Actividad autorizada». Son la misma columna.
    """
    c = re.sub(r"\s*/\s*", "/", _plano(celda))
    return c.replace("atividad", "actividad")


def leer_hoja(contenido: bytes) -> list[dict[str, Any]]:
    """Las filas de la exportación del buscador.

    El reconocimiento vio que la «hoja» es un ZIP con un XLSX dentro, y que
    `openpyxl` en modo de sólo lectura ve una fila porque la hoja no declara
    sus dimensiones: se abre en modo normal. Cabeceras como en `leer_tablas`,
    por nombre.
    """
    import openpyxl

    try:
        datos = contenido
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            interiores = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
            if interiores:
                datos = z.read(interiores[0])
        with warnings.catch_warnings():
            # «Workbook contains no default style»: la hoja del portal no trae
            # estilos, y a quien la lee le da igual.
            warnings.simplefilter("ignore", UserWarning)
            libro = openpyxl.load_workbook(io.BytesIO(datos))
    except Exception as exc:  # un formato nuevo no tumba la ingesta
        log.warning("oci: la hoja no se puede abrir", detalle=str(exc))
        return []
    filas = [
        ["" if v is None else " ".join(str(v).split()) for v in fila]
        for fila in libro.worksheets[0].iter_rows(values_only=True)
    ]
    if not filas or tuple(_cabecera(c) for c in filas[0][: len(COLUMNAS)]) != COLUMNAS:
        log.warning("oci: la hoja no trae las columnas esperadas", cabecera=filas[:1])
        return []
    salida = []
    for celdas in filas[1:]:
        if len(celdas) < len(COLUMNAS) or not celdas[0] or not celdas[4]:
            continue
        salida.append(
            {
                "nombre": celdas[0],
                "cargo": celdas[1],
                "ministerio": celdas[2],
                "fecha_cese": celdas[3],
                "actividad": celdas[4],
                "fecha_autorizacion": celdas[5],
                "curriculum": "",
            }
        )
    return salida


def leer_tablas(html: str) -> list[dict[str, Any]]:
    """Las filas de las tablas de autorizaciones de una página.

    Se reconoce la tabla por sus cabeceras, no por su posición: la página
    lleva otras tablas («Fuente de los datos»). Una tabla cuyas cabeceras no
    son éstas no se lee. El buscador pinta la misma tabla dos veces —para
    pantalla ancha y estrecha—: una fila repetida es la misma autorización.
    """
    lector = _Tablas()
    lector.feed(html)
    filas: list[dict[str, Any]] = []
    vistas: set[tuple[str, ...]] = set()
    for tabla, enlaces in zip(lector.tablas, lector.enlaces, strict=True):
        if not tabla or tuple(_cabecera(c) for c in tabla[0]) != COLUMNAS:
            continue
        for celdas, enlace in zip(tabla[1:], enlaces[1:], strict=True):
            if len(celdas) != len(COLUMNAS) or not celdas[0] or not celdas[4]:
                continue
            # El pie de la tabla repite la cabecera.
            if tuple(_cabecera(c) for c in celdas) == COLUMNAS:
                continue
            if tuple(celdas) in vistas:
                continue
            vistas.add(tuple(celdas))
            filas.append(
                {
                    "nombre": celdas[0],
                    "cargo": celdas[1],
                    "ministerio": celdas[2],
                    "fecha_cese": celdas[3],
                    "actividad": celdas[4],
                    "fecha_autorizacion": celdas[5],
                    "curriculum": enlace,
                }
            )
    return filas


class OCIConnector:
    """Lee las autorizaciones que publica el Portal de Transparencia."""

    source_id = "oci"
    extractor_version = "oci-autorizaciones/1"

    def __init__(
        self,
        cliente: httpx.Client | None = None,
        buscador: str = BUSCADOR,
        copia: Path | None = None,
        esperas: tuple[float, ...] = (10.0, 30.0),
    ):
        self._cliente = cliente
        self._buscador = buscador
        # Dónde se guarda la última exportación descargada. El 26/9/2026 el
        # buscador no respondió, la base de la ingesta empieza vacía cada
        # noche, y la web se quedó un día sin ninguna autorización. Con la
        # copia, un día sin respuesta enseña la lista del último día que la
        # hubo —con su fecha de descarga, y diciéndolo—, no un hueco.
        self._copia = copia
        self._esperas = esperas

    def _pedir(self, cliente: httpx.Client, url: str | httpx.URL) -> httpx.Response:
        """GET con reintentos: la fuente a veces tarda más de un minuto."""
        for intento, espera in enumerate((*self._esperas, None)):
            try:
                r = cliente.get(url)
                r.raise_for_status()
                return r
            except httpx.HTTPError as exc:
                if espera is None:
                    raise
                log.warning("oci: reintento", intento=intento + 1, detalle=str(exc))
                time.sleep(espera)
        raise AssertionError("inalcanzable")

    def _guardar_copia(self, raw: RawDocument) -> None:
        if self._copia is None:
            return
        self._copia.mkdir(parents=True, exist_ok=True)
        (self._copia / "exportacion.bin").write_bytes(raw.content)
        (self._copia / "exportacion.json").write_text(
            json.dumps(
                {
                    "url": raw.url,
                    "media_type": raw.media_type,
                    "retrieved_at": raw.retrieved_at.isoformat(),
                }
            ),
            encoding="utf-8",
        )

    def _desde_la_copia(self, motivo: str) -> Iterator[RawDocument]:
        """La última exportación guardada, con SU fecha de descarga, o nada."""
        if self._copia is None or not (self._copia / "exportacion.bin").exists():
            log.warning("oci: sin respuesta y sin copia guardada", motivo=motivo)
            return
        meta = json.loads((self._copia / "exportacion.json").read_text(encoding="utf-8"))
        log.error(
            "oci: la fuente no responde; se usa la copia de la última descarga",
            motivo=motivo,
            descargada=meta["retrieved_at"],
        )
        yield RawDocument(
            source_id=self.source_id,
            url=meta["url"],
            content=(self._copia / "exportacion.bin").read_bytes(),
            media_type=meta["media_type"],
            retrieved_at=datetime.fromisoformat(meta["retrieved_at"]),
            metadata={"copia": True},
        )

    def fetch(self, **_: Any) -> Iterator[RawDocument]:
        """La exportación entera del buscador, de una vez.

        Sin rango de fechas: la fuente publica la lista acumulada, y la
        idempotencia la da el hash. El exportador corta en 2000 resultados;
        si un día se pasa, se dice en el log, porque faltarían filas.
        """
        cliente = self._cliente or httpx.Client(
            timeout=60.0,
            follow_redirects=True,
            headers={"User-Agent": "Sinapsis/0.1 (proyecto abierto de transparencia)"},
        )
        propio = self._cliente is None
        try:
            try:
                r = self._pedir(cliente, self._buscador)
            except httpx.HTTPError as exc:
                log.warning("oci: el buscador no responde", detalle=str(exc))
                yield from self._desde_la_copia("el buscador no responde")
                return
            enlace = _EXPORTACION.search(r.text)
            encontrados = re.search(r"(\d[\d.]*)\s+Resultados encontrados", r.text)
            if encontrados and int(encontrados.group(1).replace(".", "")) >= 2000:
                log.warning(
                    "oci: el exportador corta en 2000 y hay más", total=encontrados.group(1)
                )
            if not enlace:
                log.warning("oci: el buscador ya no enlaza la exportación")
                yield from self._desde_la_copia("el buscador ya no enlaza la exportación")
                return
            url = httpx.URL(self._buscador).join(enlace.group(1).replace("&amp;", "&"))
            time.sleep(0.5)
            try:
                x = self._pedir(cliente, url)
            except httpx.HTTPError as exc:
                log.warning("oci: la exportación no se descarga", detalle=str(exc))
                yield from self._desde_la_copia("la exportación no se descarga")
                return
            raw = RawDocument(
                source_id=self.source_id,
                url=str(url),
                content=x.content,
                media_type=x.headers.get("content-type", "application/zip").split(";")[0],
                retrieved_at=datetime.now(UTC),
            )
            # Sólo se guarda lo que se puede leer: una exportación rota no
            # debe tapar la última buena.
            if list(self.parse(raw)):
                self._guardar_copia(raw)
            yield raw
        finally:
            if propio:
                cliente.close()

    def parse(self, raw: RawDocument) -> Iterator[ParsedRecord]:
        # La exportación (un ZIP con la hoja) o una página HTML con la tabla:
        # las dos formas que publica la fuente, y los golden tests prueban
        # las dos.
        if raw.content[:2] == b"PK":
            filas = leer_hoja(raw.content)
        else:
            filas = leer_tablas(raw.content.decode("utf-8", errors="replace"))
        for n, fila in enumerate(filas):
            yield ParsedRecord(
                raw_content_hash=raw.content_hash,
                extractor_version=self.extractor_version,
                data={**fila, "id_registro": f"{raw.url}#{n}", "url": raw.url},
            )

    def normalize(self, record: ParsedRecord) -> Normalizado | None:
        d = record.data
        cese = _fecha(d.get("fecha_cese", ""))
        autorizacion = _fecha(d.get("fecha_autorizacion", ""))
        if not d.get("nombre") or not d.get("cargo") or not d.get("actividad"):
            return None

        nombre = en_minusculas(d["nombre"])
        clave_persona = f"oci:persona:{clave(d['nombre'])}"
        clave_puesto = f"oci:puesto:{clave(d['cargo'])}"
        clave_actividad = f"oci:actividad:{clave(d['actividad'])}"
        curriculum = d.get("curriculum") or ""
        if curriculum.startswith("/"):
            curriculum = BASE + curriculum

        entidades = [
            EntidadNormalizada(
                ftm_schema="Person",
                caption=nombre,
                dedupe_key=clave_persona,
                country="es",
                # La misma marca que el BOE: la regla de §12 deja publicar a
                # quien la Oficina de Conflictos de Intereses trata como ex
                # alto cargo, y sólo en ese papel.
                properties={
                    "name": nombre,
                    "cargo_publico": True,
                    "nombreParaCruzar": nombre_para_cruzar(d["nombre"]),
                },
            ),
            EntidadNormalizada(
                ftm_schema="Position",
                caption=d["cargo"],
                dedupe_key=clave_puesto,
                country="es",
                properties={"name": d["cargo"], "ministerio": d.get("ministerio", "")},
            ),
            EntidadNormalizada(
                ftm_schema="Organization",
                caption=d["actividad"],
                dedupe_key=clave_actividad,
                country="es",
                # No es una organización identificada: es el texto de la
                # autorización, que a veces nombra una empresa y a veces dice
                # «economista por cuenta propia». Se guarda tal cual.
                properties={"name": d["actividad"], "textoDeAutorizacion": True},
            ),
        ]
        comun = {"fuente": "Oficina de Conflictos de Intereses", "url": d["url"]}
        aristas = [
            AristaNormalizada(
                ftm_schema="Occupancy",
                source_key=clave_persona,
                target_key=clave_puesto,
                # Una persona con dos autorizaciones tiene UN cese: la misma
                # clave, que es lo que lo hace idempotente.
                dedupe_key=(
                    f"oci:cese:{clave(d['nombre'])}:{clave(d['cargo'])}:{d.get('fecha_cese', '')}"
                ),
                confidence=1.0,
                end_date=cese,
                properties={
                    **comun,
                    "acto": "cese",
                    "cargo": d["cargo"],
                    "departamento": d.get("ministerio", ""),
                    **({"curriculum": curriculum} if curriculum else {}),
                },
            ),
            AristaNormalizada(
                ftm_schema="UnknownLink",
                source_key=clave_persona,
                target_key=clave_actividad,
                dedupe_key=(
                    f"oci:autorizacion:{clave(d['nombre'])}:{clave(d['actividad'])}"
                    f":{d.get('fecha_autorizacion', '')}"
                ),
                confidence=1.0,
                start_date=autorizacion,
                properties={
                    **comun,
                    "relacion": "autorizacion_actividad_privada",
                    "actividad": d["actividad"],
                    "cargoAnterior": d["cargo"],
                    **({"fechaCese": cese.isoformat()} if cese else {}),
                },
            ),
        ]
        return Normalizado(entidades=entidades, aristas=aristas)


def crear() -> OCIConnector:
    copia = os.environ.get("SINAPSIS_COPIA_OCI")
    return OCIConnector(copia=Path(copia) if copia else None)
